"""Compile the VLM result figures into VLM_results.pptx."""
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
TMP = Path("/tmp/ppt")

SLIDES = [
    ("VLM labels vs. neural features (2026-02-19)",
     [BASE / "neural_vlm_timeline.png"],
     "Lounge/sleep VLM labels above z-scored left/right neural features (10 s windows)."),
    ("Feature vs. behavior: ROC AUC", [TMP / "auc_neural_features.png"],
     "AUC strength = max(AUC, 1-AUC), windows with the camera showing the subject fully."),
    ("Feature vs. behavior: Cohen's d", [TMP / "cohens_d_neural_features.png"], "Behavior present minus absent."),
    ("Feature vs. behavior: permutation p", [TMP / "perm_p_neural_features.png"],
     "Circular time-shift permutation test; clipped at p=0.5; uncorrected."),
    ("Spectra and VLM labels, 0-40 Hz", [BASE / "VLM_spectra.png"],
     "Full 0.1 Hz-resolution spectra (ln power) time-aligned with lounge and sleep labels."),
    ("Spectra and VLM labels, 0-90 Hz", [BASE / "VLM_spectra_0-90Hz.png"], ""),
]
for metric, name in (("roc_auc", "ROC AUC strength"), ("cohens_d", "Cohen's d"), ("perm_p", "Permutation p")):
    for suffix, rng in (("", "0-40 Hz"), ("_0-90Hz", "0-90 Hz")):
        SLIDES.append((f"Frequency x behavior: {name}, {rng}",
                       [TMP / f"VLM_spectra_corr_map_{metric}{suffix}.png"],
                       "0.1 Hz bins, log power; rows = frequency, columns = behavior. "
                       + ("Uncorrected for multiple comparisons; adjacent bins are correlated."
                          if metric == "perm_p" else "")))


def add_picture(slide, path, left, top, width, height):
    with Image.open(path) as image:
        ratio = image.width / image.height
    if width / height > ratio:
        w, h = height * ratio, height
    else:
        w, h = width, width / ratio
    slide.shapes.add_picture(str(path), left + (width - w) / 2, top, w, h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

title = prs.slides.add_slide(prs.slide_layouts[0])
title.shapes.title.text = "VLM labels vs. neural activity"
title.placeholders[1].text = "TRBD001, visit 4 (2026-02-19)"

for heading, images, note in SLIDES:
    slide = prs.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6))
    box.text_frame.text = heading
    box.text_frame.paragraphs[0].font.size = Pt(26)
    box.text_frame.paragraphs[0].font.bold = True
    for path in images:
        add_picture(slide, path, Inches(0.4), Inches(0.85), Inches(12.5), Inches(5.9))
    if note:
        foot = slide.shapes.add_textbox(Inches(0.4), Inches(6.85), Inches(12.5), Inches(0.5))
        foot.text_frame.text = note
        foot.text_frame.paragraphs[0].font.size = Pt(13)

out = BASE / "VLM_results.pptx"
prs.save(out)
print("Saved", out)
